from pathlib import Path
from io import BytesIO
import warnings
import numpy as np
import pandas as pd
import streamlit as st
import matplotlib.pyplot as plt
import seaborn as sns
from sklearn.preprocessing import StandardScaler
from sklearn.decomposition import PCA
from sklearn.cluster import KMeans, AgglomerativeClustering, DBSCAN
from sklearn.metrics import silhouette_score, davies_bouldin_score, adjusted_rand_score, normalized_mutual_info_score
from sklearn.neighbors import NearestNeighbors
from scipy import stats
from scipy.cluster.hierarchy import linkage, dendrogram

st.set_page_config(page_title="Student Learning-Style Clustering", layout="wide")
sns.set_theme(style="whitegrid")
RANDOM_STATE = 42
EXPECTED_COLUMNS = ["StudyHours", "Attendance", "Resources", "Extracurricular", "Motivation", "Internet", "Gender", "Age", "LearningStyle", "OnlineCourses", "Discussions", "AssignmentCompletion", "ExamScore", "EduTech", "StressLevel", "FinalGrade"]
CLUSTER_FEATURES = ["StudyHours", "Attendance", "AssignmentCompletion", "OnlineCourses", "Discussions", "EduTech"]

st.title("Student Learning-Style Clustering Based on Learning Behavior")
st.caption("AI/ML Engineer Certification Capstone | Based on the supplied Google Colab notebook")

@st.cache_data(show_spinner=False)
def load_dataset(uploaded_bytes=None, uploaded_name=None):
    if uploaded_bytes is not None:
        return pd.read_csv(BytesIO(uploaded_bytes)), f"Uploaded CSV: {uploaded_name}"
    # First preference: the public Kaggle dataset used in the notebook.
    try:
        import kagglehub
        downloaded = Path(kagglehub.dataset_download("adilshamim8/student-performance-and-learning-style"))
        csv_files = sorted(downloaded.rglob("*.csv"))
        if csv_files:
            return pd.read_csv(csv_files[0]), f"KaggleHub: {csv_files[0].name}"
        kaggle_error = "KaggleHub did not return a CSV file."
    except Exception as exc:
        kaggle_error = str(exc)
    # Local files and Zenodo are fallbacks, matching the notebook's intent.
    for candidate in [Path("student_performance.csv"), Path("merged_dataset.csv")]:
        if candidate.exists():
            return pd.read_csv(candidate), f"Local file: {candidate.name}"
    try:
        import requests
        url = "https://zenodo.org/records/16459132/files/merged_dataset.csv?download=1"
        response = requests.get(url, timeout=60)
        response.raise_for_status()
        return pd.read_csv(BytesIO(response.content)), "Zenodo fallback"
    except Exception as exc:
        raise RuntimeError(f"Kaggle download failed ({kaggle_error}); Zenodo fallback also failed ({exc}). Upload the exact CSV used in Colab.") from exc

@st.cache_data(show_spinner="Running the notebook analysis…")
def analyse(raw_csv_bytes=None, raw_csv_name=None):
    source_df, source = load_dataset(raw_csv_bytes, raw_csv_name)
    missing = [c for c in EXPECTED_COLUMNS if c not in source_df.columns]
    if missing:
        raise ValueError("Dataset is missing required columns: " + ", ".join(missing))
    raw_n = len(source_df)
    df = source_df[EXPECTED_COLUMNS].copy()
    duplicate_count = int(df.duplicated().sum())
    df = df.drop_duplicates().reset_index(drop=True)
    X = df[CLUSTER_FEATURES].copy()
    if X.isna().sum().sum() > 0:
        X = X.fillna(X.median(numeric_only=True))
    scaler = StandardScaler()
    X_scaled = scaler.fit_transform(X)
    X_scaled_df = pd.DataFrame(X_scaled, columns=CLUSTER_FEATURES, index=df.index)
    pca = PCA(n_components=2, random_state=RANDOM_STATE)
    X_pca = pca.fit_transform(X_scaled)
    loadings = pd.DataFrame(pca.components_.T, index=CLUSTER_FEATURES, columns=["PC1", "PC2"])
    k_values = list(range(2, 9))
    inertias, silhouettes = [], []
    for k in k_values:
        model = KMeans(n_clusters=k, random_state=RANDOM_STATE, n_init=20)
        labels = model.fit_predict(X_scaled)
        inertias.append(float(model.inertia_))
        silhouettes.append(float(silhouette_score(X_scaled, labels)))
    k_results = pd.DataFrame({"k": k_values, "inertia": inertias, "silhouette": silhouettes})
    best_k = int(k_results.loc[k_results["silhouette"].idxmax(), "k"])
    best_silhouette = float(k_results["silhouette"].max())
    kmeans = KMeans(n_clusters=best_k, random_state=RANDOM_STATE, n_init=50)
    df["Cluster"] = kmeans.fit_predict(X_scaled)
    k_sil = float(silhouette_score(X_scaled, df["Cluster"]))
    k_dbi = float(davies_bouldin_score(X_scaled, df["Cluster"]))
    agg = AgglomerativeClustering(n_clusters=best_k, linkage="ward")
    df["Cluster_Agglomerative"] = agg.fit_predict(X_scaled)
    agg_sil = float(silhouette_score(X_scaled, df["Cluster_Agglomerative"]))
    agg_dbi = float(davies_bouldin_score(X_scaled, df["Cluster_Agglomerative"]))
    min_samples = 5
    nn = NearestNeighbors(n_neighbors=min_samples).fit(X_scaled)
    distances, _ = nn.kneighbors(X_scaled)
    kdist = np.sort(distances[:, -1])
    candidate_eps = np.unique(np.round(np.quantile(kdist, [0.80, 0.85, 0.90, 0.92, 0.94, 0.95, 0.96, 0.97, 0.98]), 3))
    db_rows = []
    for eps in candidate_eps:
        labels = DBSCAN(eps=float(eps), min_samples=min_samples).fit_predict(X_scaled)
        n_clusters = len(set(labels)) - (1 if -1 in labels else 0)
        noise_pct = 100 * np.mean(labels == -1)
        if n_clusters >= 2 and np.sum(labels != -1) > n_clusters:
            mask = labels != -1
            sil = float(silhouette_score(X_scaled[mask], labels[mask]))
            dbi = float(davies_bouldin_score(X_scaled[mask], labels[mask]))
        else:
            sil, dbi = np.nan, np.nan
        db_rows.append({"eps": float(eps), "clusters": n_clusters, "noise_percent": noise_pct, "silhouette": sil, "davies_bouldin": dbi})
    db_results = pd.DataFrame(db_rows)
    valid_db = db_results.dropna(subset=["silhouette"]).copy()
    if len(valid_db):
        preferred = valid_db[valid_db["noise_percent"] <= 20]
        db_choice = preferred if len(preferred) else valid_db
        best_eps = float(db_choice.sort_values(["silhouette", "noise_percent"], ascending=[False, True]).iloc[0]["eps"])
        df["Cluster_DBSCAN"] = DBSCAN(eps=best_eps, min_samples=min_samples).fit_predict(X_scaled)
        db_mask = df["Cluster_DBSCAN"] != -1
        db_clusters = len(set(df["Cluster_DBSCAN"])) - (1 if -1 in df["Cluster_DBSCAN"].values else 0)
        if db_clusters >= 2 and db_mask.sum() > db_clusters:
            db_sil = float(silhouette_score(X_scaled[db_mask], df.loc[db_mask, "Cluster_DBSCAN"]))
            db_dbi = float(davies_bouldin_score(X_scaled[db_mask], df.loc[db_mask, "Cluster_DBSCAN"]))
        else:
            db_sil, db_dbi = np.nan, np.nan
    else:
        best_eps, db_clusters, db_sil, db_dbi = np.nan, 0, np.nan, np.nan
        df["Cluster_DBSCAN"] = -1
    comparison = pd.DataFrame({"Model": ["K-Means", "Agglomerative", "DBSCAN"], "Silhouette": [k_sil, agg_sil, db_sil], "Davies_Bouldin": [k_dbi, agg_dbi, db_dbi]})
    best_model = comparison.loc[comparison["Silhouette"].idxmax(), "Model"] if comparison["Silhouette"].notna().any() else "K-Means"
    profile_cols = CLUSTER_FEATURES + ["Motivation", "StressLevel", "Resources", "Internet", "Extracurricular", "ExamScore", "FinalGrade"]
    cluster_profile = df.groupby("Cluster")[profile_cols].mean().round(2)
    centroid_profile = pd.DataFrame(kmeans.cluster_centers_, columns=CLUSTER_FEATURES).round(2)
    centroid_profile.index.name = "Cluster"
    crosstab = pd.crosstab(df["Cluster"], df["LearningStyle"], normalize="index").mul(100).round(1)
    ari = float(adjusted_rand_score(df["LearningStyle"], df["Cluster"]))
    nmi = float(normalized_mutual_info_score(df["LearningStyle"], df["Cluster"]))
    outcome_profile = df.groupby("Cluster")[["ExamScore", "FinalGrade"]].agg(["mean", "median", "std"]).round(2)
    groups = [g["ExamScore"].values for _, g in df.groupby("Cluster")]
    anova = stats.f_oneway(*groups)
    grand_mean = df["ExamScore"].mean()
    ss_between = sum(len(g) * (g.mean() - grand_mean) ** 2 for g in groups)
    ss_total = ((df["ExamScore"] - grand_mean) ** 2).sum()
    eta_squared = float(ss_between / ss_total) if ss_total else np.nan
    context_profile = df.groupby("Cluster")[["Motivation", "StressLevel", "Resources", "Internet", "Extracurricular"]].mean().round(2)
    global_means = df[CLUSTER_FEATURES].mean()
    def describe_cluster(cid, centroid_row, top_n=2):
        ordered = centroid_row.sort_values()
        return f"Cluster {cid}: relatively high {', '.join(ordered.tail(top_n).index)}; relatively low {', '.join(ordered.head(top_n).index)}."
    def make_recommendation(cid):
        row = cluster_profile.loc[cid, CLUSTER_FEATURES]
        recs = []
        checks = [("Attendance", "attendance support and reminders"), ("AssignmentCompletion", "assignment planning and deadline support"), ("StudyHours", "structured study scheduling"), ("OnlineCourses", "guided use of supplementary online learning"), ("Discussions", "peer discussion or collaborative learning opportunities"), ("EduTech", "orientation to educational technology tools")]
        for col, rec in checks:
            if row[col] < global_means[col]:
                recs.append(rec)
        return "; ".join(recs[:3]) if recs else "enrichment activities and advanced self-directed resources"
    recommendation_table = pd.DataFrame({
        "Cluster": sorted(df["Cluster"].unique()),
        "Observed_behavior": [describe_cluster(cid, centroid_profile.loc[cid]) for cid in sorted(df["Cluster"].unique())],
        "Recommended_support": [make_recommendation(cid) for cid in sorted(df["Cluster"].unique())]
    })
    sample_n = min(1000, len(X_scaled))
    rng = np.random.default_rng(RANDOM_STATE)
    sample_idx = rng.choice(len(X_scaled), size=sample_n, replace=False)
    linked = linkage(X_scaled[sample_idx], method="ward")
    # Streamlit caches this function result; nested helper functions are not safely serializable.
    analysis_results = locals().copy()
    return {key: value for key, value in analysis_results.items() if not callable(value)}

uploaded = st.sidebar.file_uploader("Upload the same CSV used in Colab (optional)", type=["csv"])
st.sidebar.caption("Without an upload, the app tries Kaggle first, then local CSV files, then Zenodo.")
try:
    result = analyse(uploaded.getvalue() if uploaded else None, uploaded.name if uploaded else None)
except Exception as e:
    st.error(f"Could not load or analyse the dataset: {e}")
    st.info("If Kaggle access is unavailable, upload the exact CSV used in your Colab notebook.")
    st.stop()

df = result["df"]
data = df[EXPECTED_COLUMNS].copy()
source = result["source"]
st.sidebar.success(f"Data source: {source}")
page = st.sidebar.radio("Choose a section", [
    "Project overview", "Data quality audit", "Data explorer", "Exploratory data analysis",
    "Feature engineering and PCA", "K-Means model selection", "Cluster discovery", "Model comparison",
    "Cluster profiles and recommendations", "Learning-style validation", "Academic outcomes", "Executive summary"
])

if page == "Project overview":
    st.subheader("Project objective")
    st.write("Discover student learning-behavior profiles using unsupervised machine learning. LearningStyle, ExamScore and FinalGrade are excluded from the clustering inputs and examined afterward.")
    st.markdown("**Six clustering features**")
    st.code(", ".join(CLUSTER_FEATURES))
    c1,c2,c3=st.columns(3)
    c1.metric("Raw records", f'{result["raw_n"]:,}')
    c2.metric("Exact duplicates removed", f'{result["duplicate_count"]:,}')
    c3.metric("Analytical records", f'{len(df):,}')
    st.markdown("**Notebook workflow:** data audit → behavioral feature selection → standardization → PCA → K-Means selection → K-Means, Agglomerative and DBSCAN comparison → interpretation and outcome analysis.")
elif page == "Data quality audit":
    st.subheader("Data quality audit")
    st.write("The notebook removes exact duplicate rows only and does not automatically remove outliers.")
    st.write("Initial raw rows:", result["raw_n"], "| Exact duplicate rows:", result["duplicate_count"], "| Rows after duplicate removal:", len(data))
    st.markdown("**Data types (original 16 columns only)**")
    st.dataframe(data.dtypes.astype(str).rename("dtype").to_frame(), use_container_width=True)
    st.markdown("**Missing values (original 16 columns only)**")
    st.dataframe(data.isna().sum().rename("missing").to_frame(), use_container_width=True)
    st.markdown("**Numeric summary (original 16 columns only)**")
    st.dataframe(data.describe().T, use_container_width=True)
    checks = {"Age": (18,30), "Attendance": (0,100), "AssignmentCompletion": (0,100), "ExamScore": (0,100)}
    st.markdown("**Basic range checks**")
    st.dataframe(pd.DataFrame([{"Column": c, "Allowed minimum": lo, "Allowed maximum": hi, "Values outside range": int(((data[c] < lo) | (data[c] > hi)).sum())} for c,(lo,hi) in checks.items()]), use_container_width=True)
elif page == "Data explorer":
    st.subheader("Cleaned source dataset preview")
    st.dataframe(data.head(100), use_container_width=True)
    st.markdown("**Missing values**")
    st.dataframe(data.isna().sum().rename("Missing values").to_frame(), use_container_width=True)
    st.markdown("**Data types**")
    st.dataframe(data.dtypes.astype(str).rename("Data type").to_frame(), use_container_width=True)
    st.download_button("Download cleaned source dataset", data.to_csv(index=False).encode(), "student_learning_style_cleaned_source.csv", "text/csv")
elif page == "Exploratory data analysis":
    st.subheader("Distribution of core behavioral variables")
    fig, axes = plt.subplots(2,3,figsize=(15,8))
    for ax,col in zip(axes.flat, CLUSTER_FEATURES):
        sns.histplot(data[col], kde=True, ax=ax); ax.set_title(col)
    st.pyplot(fig); plt.close(fig)
    fig, ax = plt.subplots(figsize=(9,7))
    sns.heatmap(data[CLUSTER_FEATURES].corr(), annot=True, fmt=".2f", cmap="vlag", center=0, ax=ax)
    ax.set_title("Correlation Matrix: Core Learning-Behavior Variables")
    st.pyplot(fig); plt.close(fig)
    fig, axes = plt.subplots(1,3,figsize=(16,5))
    sns.countplot(data=data,x="LearningStyle",ax=axes[0]); axes[0].set_title("Reported Learning-Style Codes")
    sns.histplot(data["ExamScore"],kde=True,ax=axes[1]); axes[1].set_title("Exam Score Distribution")
    sns.countplot(data=data,x="FinalGrade",ax=axes[2]); axes[2].set_title("Final Grade Codes")
    st.pyplot(fig); plt.close(fig)
elif page == "Feature engineering and PCA":
    st.subheader("Feature selection")
    st.write("Only these six variables are used for clustering:")
    st.code(", ".join(CLUSTER_FEATURES))
    st.markdown("**Missing values in clustering features**")
    st.dataframe(data[CLUSTER_FEATURES].isna().sum().rename("missing").to_frame(), use_container_width=True)
    st.markdown("**Clustering feature summary**")
    st.dataframe(data[CLUSTER_FEATURES].describe().T, use_container_width=True)
    st.markdown("**Standardized feature summary**")
    st.dataframe(result["X_scaled_df"].describe().T, use_container_width=True)
    st.write("PC1 explained variance:", round(result["pca"].explained_variance_ratio_[0],4))
    st.write("PC2 explained variance:", round(result["pca"].explained_variance_ratio_[1],4))
    st.write("Total 2D explained variance:", round(result["pca"].explained_variance_ratio_.sum(),4))
    fig,ax=plt.subplots(figsize=(8,6)); ax.scatter(result["X_pca"][:,0],result["X_pca"][:,1],alpha=.35,s=18); ax.set(xlabel="PC1",ylabel="PC2",title="PCA Projection of Student Learning Behavior"); st.pyplot(fig); plt.close(fig)
    st.markdown("**PCA loadings**"); st.dataframe(result["loadings"].sort_values("PC1",key=np.abs,ascending=False),use_container_width=True)
elif page == "K-Means model selection":
    st.subheader("Elbow and silhouette analysis")
    st.dataframe(result["k_results"], use_container_width=True)
    fig,axes=plt.subplots(1,2,figsize=(13,5))
    axes[0].plot(result["k_results"]["k"],result["k_results"]["inertia"],marker="o"); axes[0].axvline(result["best_k"],linestyle="--",alpha=.7); axes[0].set(title="Elbow Method",xlabel="Number of clusters (k)",ylabel="Inertia")
    axes[1].plot(result["k_results"]["k"],result["k_results"]["silhouette"],marker="o"); axes[1].axvline(result["best_k"],linestyle="--",alpha=.7); axes[1].set(title="Silhouette Scores",xlabel="Number of clusters (k)",ylabel="Silhouette score")
    st.pyplot(fig); plt.close(fig)
    st.success(f'Selected K-Means k = {result["best_k"]}; best silhouette = {result["best_silhouette"]:.3f}')
elif page == "Cluster discovery":
    st.subheader("Cluster discovery")
    st.write(f'K-Means silhouette: {result["k_sil"]:.3f} | Davies-Bouldin: {result["k_dbi"]:.3f}')
    st.write(f'Agglomerative silhouette: {result["agg_sil"]:.3f} | Davies-Bouldin: {result["agg_dbi"]:.3f}')
    st.markdown("**K-Means cluster sizes**")
    sizes=df["Cluster"].value_counts().sort_index().rename("count").to_frame(); sizes["percentage"]=100*sizes["count"]/len(df); st.dataframe(sizes,use_container_width=True)
    fig,ax=plt.subplots(figsize=(9,7)); sns.scatterplot(x=result["X_pca"][:,0],y=result["X_pca"][:,1],hue=df["Cluster"].astype(str),alpha=.55,s=35,ax=ax); ax.set(title=f'K-Means Behavioral Clusters (k={result["best_k"]})',xlabel="PC1",ylabel="PC2"); st.pyplot(fig); plt.close(fig)
    fig,ax=plt.subplots(figsize=(12,6)); dendrogram(result["linked"],truncate_mode="lastp",p=30,no_labels=True,ax=ax); ax.set(title="Hierarchical Clustering Dendrogram (sample for readability)",xlabel="Merged observations / groups",ylabel="Ward distance"); st.pyplot(fig); plt.close(fig)
    st.markdown("**DBSCAN parameter search**"); st.dataframe(result["db_results"],use_container_width=True)
    st.write("Selected DBSCAN eps:", result["best_eps"])
    st.write("DBSCAN clusters:", result["db_clusters"])
    st.write("DBSCAN silhouette excluding noise:", result["db_sil"] if pd.notna(result["db_sil"]) else "not available")
    st.markdown("**Cluster assignments**")
    st.dataframe(df[["Cluster","Cluster_Agglomerative","Cluster_DBSCAN"]].head(100),use_container_width=True)
elif page == "Model comparison":
    st.subheader("Comparison of clustering algorithms")
    st.dataframe(result["comparison"],use_container_width=True)
    st.write("Best internal-silhouette model:", result["best_model"])
elif page == "Cluster profiles and recommendations":
    st.subheader("Cluster profiles")
    st.dataframe(result["cluster_profile"],use_container_width=True)
    st.markdown("**Standardized cluster centroids**")
    st.dataframe(result["centroid_profile"],use_container_width=True)
    fig,ax=plt.subplots(figsize=(11,6)); sns.heatmap(result["centroid_profile"],annot=True,fmt=".2f",cmap="vlag",center=0,ax=ax); ax.set_title("Standardized Behavioral Cluster Centroids"); st.pyplot(fig); plt.close(fig)
    size_plot=df["Cluster"].value_counts().sort_index().rename_axis("Cluster").reset_index(name="Students")
    fig,ax=plt.subplots(figsize=(8,5)); sns.barplot(data=size_plot,x="Cluster",y="Students",ax=ax); ax.set_title("Number of Students in Each Behavioral Cluster"); st.pyplot(fig); plt.close(fig)
    st.markdown("**Empirical descriptions and support recommendations**")
    st.dataframe(result["recommendation_table"],use_container_width=True)
    st.markdown("**Contextual profile**"); st.dataframe(result["context_profile"],use_container_width=True)
elif page == "Learning-style validation":
    st.subheader("Behavioral clusters versus reported LearningStyle")
    st.dataframe(result["crosstab"],use_container_width=True)
    c1,c2=st.columns(2); c1.metric("Adjusted Rand Index",f'{result["ari"]:.3f}'); c2.metric("Normalized Mutual Information",f'{result["nmi"]:.3f}')
    fig,ax=plt.subplots(figsize=(9,6)); sns.heatmap(result["crosstab"],annot=True,fmt=".1f",cmap="Blues",cbar_kws={"label":"Percent within cluster"},ax=ax); ax.set(title="Learning-Style Composition of Behavioral Clusters",xlabel="LearningStyle code",ylabel="Behavioral cluster"); st.pyplot(fig); plt.close(fig)
    st.caption("Reported LearningStyle is evaluated after clustering; it is not used to create the clusters.")
elif page == "Academic outcomes":
    st.subheader("Academic outcomes by behavioral cluster")
    st.dataframe(result["outcome_profile"],use_container_width=True)
    fig,ax=plt.subplots(figsize=(10,6)); sns.boxplot(data=df,x="Cluster",y="ExamScore",ax=ax); ax.set_title("Exam Score by Behavioral Cluster"); st.pyplot(fig); plt.close(fig)
    st.write(f'ANOVA F-statistic: {result["anova"].statistic:.3f}')
    st.write(f'ANOVA p-value: {result["anova"].pvalue:.6g}')
    st.write(f'Eta-squared: {result["eta_squared"]:.3f}')
    st.write("Interpretation: mean ExamScore differs significantly across at least some behavioral clusters." if result["anova"].pvalue<.05 else "Interpretation: no statistically significant difference in mean ExamScore was detected across clusters.")
elif page == "Executive summary":
    st.subheader("Capstone executive summary")
    summary = {"Raw records available": result["raw_n"], "Exact duplicates removed": result["duplicate_count"], "Final analytical records": len(df), "Behavioral features used": len(CLUSTER_FEATURES), "Selected K-Means clusters": result["best_k"], "K-Means silhouette": round(result["k_sil"],3), "K-Means Davies-Bouldin": round(result["k_dbi"],3), "Agglomerative silhouette": round(result["agg_sil"],3), "DBSCAN silhouette": round(result["db_sil"],3) if pd.notna(result["db_sil"]) else "not available", "LearningStyle ARI": round(result["ari"],3), "LearningStyle NMI": round(result["nmi"],3), "ExamScore ANOVA p-value": f'{result["anova"].pvalue:.6g}', "ExamScore eta-squared": round(result["eta_squared"],3)}
    st.dataframe(pd.DataFrame({"Metric":summary.keys(),"Value":summary.values()}),hide_index=True,use_container_width=True)
    st.markdown("**Cluster profiles and recommended support**")
    st.dataframe(result["recommendation_table"],use_container_width=True)
    st.download_button("Download analysed dataset (includes cluster labels)",df.to_csv(index=False).encode(),"student_learning_style_clustered_dataset.csv","text/csv")
    st.markdown("**Capstone PowerPoint**")
    from pptx import Presentation
    from pptx.util import Inches, Pt
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    def add_bullet_slide(title, bullets):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = title
        frame = slide.placeholders[1].text_frame
        frame.clear()
        for i, bullet in enumerate(bullets):
            paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            paragraph.text = str(bullet)
            paragraph.level = 0
            paragraph.font.size = Pt(18)
    title_slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    title_slide.shapes.title.text = "Student Learning-Style Clustering Based on Learning Behavior"
    title_slide.placeholders[1].text = "AI/ML Engineer Certification Capstone"
    add_bullet_slide("1. Problem and Objective", [
        "Discover student learning-behavior profiles using unsupervised machine learning.",
        "K-Means is validated against hierarchical clustering and DBSCAN.",
        "LearningStyle is evaluated after clustering, not used to create clusters.",
        "Academic performance is held out from clustering and examined afterward."
    ])
    add_bullet_slide("2. Dataset", [
        f"Source: {source}.",
        f"Raw records: {result['raw_n']:,}.",
        f"Exact duplicates removed: {result['duplicate_count']:,}.",
        f"Final analytical records: {len(df):,}."
    ])
    add_bullet_slide("3. Behavioral Feature Strategy", CLUSTER_FEATURES + [
        "LearningStyle, ExamScore and FinalGrade were excluded from clustering."
    ])
    add_bullet_slide("4. Model Selection", [
        f"Selected K-Means k: {result['best_k']}.",
        f"K-Means silhouette: {result['k_sil']:.3f}.",
        f"K-Means Davies-Bouldin index: {result['k_dbi']:.3f}.",
        f"Agglomerative silhouette: {result['agg_sil']:.3f}.",
        f"DBSCAN silhouette: {result['db_sil']:.3f}" if pd.notna(result['db_sil']) else "DBSCAN did not produce a valid multi-cluster solution."
    ])
    add_bullet_slide("5. Behavioral Cluster Profiles", [
        f"Cluster {int(row['Cluster'])}: {row['Observed_behavior']}"
        for _, row in result["recommendation_table"].iterrows()
    ])
    add_bullet_slide("6. Learning-Style Validation", [
        f"Adjusted Rand Index: {result['ari']:.3f}.",
        f"Normalized Mutual Information: {result['nmi']:.3f}.",
        "Reported learning style is an external comparison, not a clustering input."
    ])
    add_bullet_slide("7. Academic Outcome Analysis", [
        f"ExamScore ANOVA p-value: {result['anova'].pvalue:.6g}.",
        f"Eta-squared: {result['eta_squared']:.3f}.",
        "ExamScore was excluded from clustering and analyzed afterward."
    ])
    add_bullet_slide("8. Recommended Educational Support", [
        f"Cluster {int(row['Cluster'])}: {row['Recommended_support']}"
        for _, row in result["recommendation_table"].iterrows()
    ])
    add_bullet_slide("9. Limitations", [
        "Secondary dataset with incomplete documentation of recruitment and sampling.",
        "Behavioral variables are structured proxies, not a live LMS event log.",
        "Clusters are not permanent student identities.",
        "Associations do not establish causation."
    ])
    add_bullet_slide("10. Conclusion", [
        f"The pipeline identified {result['best_k']} behavioral segments using K-Means.",
        "Behavioral segmentation can support differentiated educational interventions.",
        "LearningStyle was used as an external comparison.",
        "Profiles convert learning-behavior variables into interpretable student segments."
    ])
    pptx_buffer = BytesIO()
    presentation.save(pptx_buffer)
    st.download_button(
        "Download capstone PowerPoint",
        data=pptx_buffer.getvalue(),
        file_name="student_learning_style_clustering_capstone.pptx",
        mime="application/vnd.openxmlformats-officedocument.presentationml.presentation"
    )
